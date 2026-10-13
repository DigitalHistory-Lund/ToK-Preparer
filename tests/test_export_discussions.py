"""Tests for src.export_discussions: per-discussion pdf/docx/jsonl export."""

import io
import json
import re
import sqlite3
import unittest

from src.export_discussions import (
    DOI_URL,
    PROJECT_URL,
    project_version,
    Discussion,
    Utterance,
    discussion_record,
    display_name,
    docx_bytes,
    file_stem,
    load_discussions,
    render_docx,
    render_pdf,
    split_runs,
    utterance_records,
)


SCHEMA = """
CREATE TABLE person (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    gender TEXT,
    party TEXT,
    UNIQUE(name, gender, party)
);
CREATE TABLE utterance (
    id TEXT PRIMARY KEY,
    content TEXT,
    prev TEXT,
    next TEXT,
    person_id INTEGER,
    year INTEGER,
    date INTEGER,
    kammare INTEGER,
    kvinna_1 BOOLEAN,
    kvinna_2 BOOLEAN,
    kvinna_3 BOOLEAN,
    discussion_id INTEGER,
    FOREIGN KEY (person_id) REFERENCES person(id)
);
"""

NAMES = {"i-aaa": "anna lindhagen", "i-bbb": "carl af wirsén"}


def _fixture_conn() -> sqlite3.Connection:
    """Two discussions, inserted deliberately out of chain order.

    Discussion 7 (chamber 1, 1922-05-27): a -> b -> c where b is a
    non-kvinna interjection. Discussion 9 (chamber 2, same date): x -> y.
    Chain links also reach outside the discussions (``pre``/``post``)
    so the loader must restrict the walk to the group's own members.
    """
    conn = sqlite3.connect(":memory:")
    conn.executescript(SCHEMA)
    conn.executemany(
        "INSERT INTO person (id, name, gender, party) VALUES (?, ?, ?, ?)",
        [
            (1, "i-aaa", "woman", "Socialdemokraterna"),
            (2, "i-bbb", "man", None),
            (3, "i-zzz", None, None),
        ],
    )
    rows = [
        # id, content, prev, next, person_id, year, date, kammare, k1, k2, k3, disc
        ("c", "Sista ordet om kvinnan¹.", "b", "post", 1, 1922, 19220527, 1, 1, 0, 0, 7),
        ("a", "Herr talman! Hon² talade -- ja.", "pre", "b", 2, 1922, 19220527, 1, 0, 1, 0, 7),
        ("b", "Ett inlägg utan träff.", "a", "c", 3, 1922, 19220527, 1, 0, 0, 0, 7),
        ("y", "Andra kammaren, lärarinnor³.", "x", None, 1, 1922, 19220527, 2, 0, 0, 1, 9),
        ("x", "Kvinnor¹ i andra kammaren.", None, "y", 2, 1922, 19220527, 2, 1, 0, 0, 9),
        ("pre", "Före diskussionen.", None, "a", 3, 1922, 19220527, 1, 0, 0, 0, None),
        ("post", "Efter diskussionen.", "c", None, 3, 1922, 19220527, 1, 0, 0, 0, None),
    ]
    conn.executemany("INSERT INTO utterance VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    conn.commit()
    return conn


def _sample_discussion(n_utterances: int = 2, content: str | None = None) -> Discussion:
    content = content or "Herr talman! Hon² talade om kvinnor¹ -- ja."
    utts = [
        Utterance(
            id=f"u{i}",
            order=i + 1,
            speaker_id="i-aaa",
            speaker_name="Anna Lindhagen",
            gender="woman",
            party="Socialdemokraterna",
            kvinna_1=True,
            kvinna_2=True,
            kvinna_3=False,
            content=content,
        )
        for i in range(n_utterances)
    ]
    return Discussion(
        id=423,
        year=1922,
        date="1922-05-27",
        chamber=1,
        chamber_name="Första kammaren",
        utterances=utts,
    )


class LoadDiscussions(unittest.TestCase):
    def setUp(self):
        self.conn = _fixture_conn()
        self.addCleanup(self.conn.close)

    def test_orders_by_chain_not_insertion(self):
        discs = load_discussions(self.conn, NAMES)
        by_id = {d.id: d for d in discs}
        self.assertEqual(sorted(by_id), [7, 9])
        self.assertEqual([u.id for u in by_id[7].utterances], ["a", "b", "c"])
        self.assertEqual([u.order for u in by_id[7].utterances], [1, 2, 3])
        self.assertEqual([u.id for u in by_id[9].utterances], ["x", "y"])

    def test_metadata_and_names(self):
        discs = {d.id: d for d in load_discussions(self.conn, NAMES)}
        d7 = discs[7]
        self.assertEqual(d7.year, 1922)
        self.assertEqual(d7.date, "1922-05-27")
        self.assertEqual(d7.chamber, 1)
        self.assertEqual(d7.chamber_name, "Första kammaren")
        self.assertEqual(discs[9].chamber_name, "Andra kammaren")
        a, b, c = d7.utterances
        self.assertEqual(a.speaker_id, "i-bbb")
        self.assertEqual(a.speaker_name, "Carl af Wirsén")
        self.assertEqual(a.gender, "man")
        self.assertIsNone(a.party)
        self.assertEqual(b.speaker_name, "i-zzz")  # unmapped -> raw id
        self.assertEqual(c.party, "Socialdemokraterna")
        self.assertTrue(c.kvinna_1)
        self.assertFalse(c.kvinna_2)
        self.assertFalse(b.kvinna_1)

    def test_broken_chain_raises(self):
        # Detach "c" from the chain: b no longer points at it.
        self.conn.execute("UPDATE utterance SET next = NULL WHERE id = 'b'")
        with self.assertRaises(ValueError):
            load_discussions(self.conn, NAMES)


class SplitRuns(unittest.TestCase):
    def test_marks_annotated_words_bold_and_keeps_rest(self):
        runs = split_runs("Hon¹ talade om kvinnor¹ -- ja.")
        self.assertEqual(
            runs,
            [("Hon¹", True), (" talade om ", False), ("kvinnor¹", True), (" -- ja.", False)],
        )

    def test_no_markers_gives_single_plain_run(self):
        self.assertEqual(split_runs("Inget här."), [("Inget här.", False)])

    def test_multiple_markers_on_one_word(self):
        self.assertEqual(split_runs("fru²³!"), [("fru²³", True), ("!", False)])


class ProjectVersion(unittest.TestCase):
    def test_reads_version_from_pyproject(self):
        import tomllib
        from src.settings import root

        with open(root / "pyproject.toml", "rb") as fh:
            expected = tomllib.load(fh)["project"]["version"]
        self.assertEqual(project_version(), expected)
        self.assertRegex(project_version(), r"^\d{4}\.\d{2}\.\d{2}$")


class DisplayName(unittest.TestCase):
    def test_title_cases_and_keeps_particles(self):
        self.assertEqual(display_name("oskar strutz"), "Oskar Strutz")
        self.assertEqual(display_name("carl af wirsén"), "Carl af Wirsén")
        self.assertEqual(display_name("hugo von hamilton"), "Hugo von Hamilton")


class Records(unittest.TestCase):
    def test_discussion_record_shape(self):
        rec = discussion_record(_sample_discussion(3))
        self.assertEqual(rec["discussion_id"], 423)
        self.assertEqual(rec["date"], "1922-05-27")
        self.assertEqual(rec["chamber"], 1)
        self.assertEqual(rec["chamber_name"], "Första kammaren")
        self.assertEqual(rec["n_utterances"], 3)
        self.assertEqual(len(rec["utterances"]), 3)
        u = rec["utterances"][0]
        self.assertEqual(
            sorted(u),
            sorted(
                [
                    "order", "id", "speaker_id", "speaker_name", "gender", "party",
                    "kvinna_1", "kvinna_2", "kvinna_3", "content",
                ]
            ),
        )
        json.dumps(rec)  # serialisable

    def test_utterance_records_carry_discussion_keys(self):
        recs = utterance_records(_sample_discussion(2))
        self.assertEqual(len(recs), 2)
        self.assertEqual(recs[1]["order"], 2)
        self.assertEqual(recs[1]["discussion_id"], 423)
        self.assertEqual(recs[1]["date"], "1922-05-27")
        self.assertEqual(recs[1]["chamber"], 1)

    def test_records_carry_source_urls_and_version(self):
        expected = {
            "project_url": PROJECT_URL,
            "doi_url": DOI_URL,
            "version": project_version(),
        }
        rec = discussion_record(_sample_discussion(1))
        self.assertEqual(rec["source"], expected)
        for u in utterance_records(_sample_discussion(2)):
            self.assertEqual(u["source"], expected)
        self.assertEqual(PROJECT_URL, "https://github.com/DigitalHistory-Lund/ToK-Preparer")
        self.assertEqual(DOI_URL, "https://doi.org/10.5281/zenodo.18211787")

    def test_file_stem(self):
        self.assertEqual(file_stem(_sample_discussion()), "1922-05-27_fk_d0423")
        d2 = _sample_discussion()
        d2.chamber = 2
        self.assertEqual(file_stem(d2), "1922-05-27_ak_d0423")


class RenderDocx(unittest.TestCase):
    def test_contains_heading_and_bold_match(self):
        doc = render_docx(_sample_discussion(1))
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any("1. Anna Lindhagen (Socialdemokraterna, woman)" in t for t in texts))
        bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
        self.assertIn("Hon²", bold_runs)
        self.assertIn("kvinnor¹", bold_runs)
        plain = [r.text for p in doc.paragraphs for r in p.runs if not r.bold]
        self.assertTrue(any("-- ja." in t for t in plain))

    def test_has_hyperlinks_to_project_and_doi(self):
        doc = render_docx(_sample_discussion(1))
        targets = {rel.target_ref for rel in doc.part.rels.values() if rel.is_external}
        self.assertIn(PROJECT_URL, targets)
        self.assertIn(DOI_URL, targets)
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn(PROJECT_URL, texts)
        self.assertIn(DOI_URL, texts)

    def test_header_shows_version(self):
        doc = render_docx(_sample_discussion(1))
        texts = [p.text for p in doc.paragraphs]
        self.assertIn(f"Version: {project_version()}", texts)

    def test_deterministic_bytes(self):
        self.assertEqual(docx_bytes(_sample_discussion(2)), docx_bytes(_sample_discussion(2)))

    def test_docx_bytes_is_valid_docx(self):
        from docx import Document

        doc = Document(io.BytesIO(docx_bytes(_sample_discussion(1))))
        self.assertTrue(any("Diskussion 423" in p.text for p in doc.paragraphs))


class RenderPdf(unittest.TestCase):
    def test_returns_pdf_bytes(self):
        out = render_pdf(_sample_discussion(1))
        self.assertTrue(out.startswith(b"%PDF"))

    def test_has_links_to_project_and_doi(self):
        out = render_pdf(_sample_discussion(1))
        self.assertIn(PROJECT_URL.encode(), out)
        self.assertIn(DOI_URL.encode(), out)
        self.assertIn(b"/URI", out)

    def test_header_shows_version(self):
        out = render_pdf(_sample_discussion(1))
        # fpdf2 subsets fonts and encodes text as glyph ids, so check the
        # version via the document metadata instead of the page stream.
        self.assertIn(f"/Subject ({project_version()}".encode(), out)

    def test_long_discussion_spans_pages(self):
        long_text = ("Herr talman! Hon² talade länge om kvinnor¹ — ”mycket”. " * 40)
        out = render_pdf(_sample_discussion(6, content=long_text))
        self.assertGreater(len(re.findall(rb"/Type ?/Page[^s]", out)), 1)

    def test_deterministic_bytes(self):
        self.assertEqual(render_pdf(_sample_discussion(2)), render_pdf(_sample_discussion(2)))


if __name__ == "__main__":
    unittest.main()
