"""Export every ``discussion_id`` arc as reference files (pdf, docx, jsonl).

Pipeline step 5. Reads only shipped artefacts — the per-year
``ToK_data_YYYY.sqlite3.gz`` files and ``speaker_totals.csv.gz`` (for the
``who -> name`` map) — so it runs from a plain checkout without ``data/``.

Output layout (``settings.discussions_dir``)::

    discussions/
      discussions.jsonl.gz          one line per discussion, utterances nested
      1922/
        1922-05-27_fk_d0423.pdf
        1922-05-27_fk_d0423.docx
        1922-05-27_fk_d0423.jsonl   one line per utterance

Output is byte-deterministic: document metadata dates are pinned, the gzip
header carries ``mtime=0`` and docx zip entries get a fixed timestamp, so
re-running the step on unchanged input leaves ``git status`` clean.
"""

from __future__ import annotations

import argparse
import contextlib
import gzip
import io
import json
import re
import sqlite3
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Iterator

from docx import Document
from docx.shared import Pt
from fpdf import FPDF
from tqdm import tqdm

from . import settings

CHAMBER_NAMES = {1: "Första kammaren", 2: "Andra kammaren"}
CHAMBER_CODES = {1: "fk", 2: "ak"}

# Labels used in the rendered documents. Swedish, because the readers of
# the reference files are the project's historians; jsonl keys stay English.
LABELS = {
    "title": "Diskussion {id}",
    "date": "Datum",
    "chamber": "Kammare",
    "n_utterances": "Antal yttranden",
    "source": "Källa",
    "legend": "Markerade ord: ¹ kvinna-ord, ² släktskaps- och tilltalsord, ³ yrkestitlar.",
    "page": "Sida {n}",
}
SOURCE_LINE = (
    "Tal om Kvinnor: Annotated Corpus (doi:10.5281/zenodo.18211787), "
    "byggd på The Swedish Parliament Corpus (Yrjänäinen m.fl. 2024)."
)

# Pinned so that unchanged input yields identical bytes across runs.
FIXED_DATETIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)

FONT_DIR = Path(__file__).resolve().parent / "fonts"
FONT_FAMILY = "DejaVu"

_RUN_RE = re.compile(r"\w+[¹²³]+")
_NAME_PARTICLES = frozenset({"af", "von", "de", "van", "der", "zu"})


@dataclass
class Utterance:
    id: str
    order: int
    speaker_id: str
    speaker_name: str
    gender: str | None
    party: str | None
    kvinna_1: bool
    kvinna_2: bool
    kvinna_3: bool
    content: str


@dataclass
class Discussion:
    id: int
    year: int
    date: str
    chamber: int
    chamber_name: str
    utterances: list[Utterance] = field(default_factory=list)


# --------------------------------------------------------------------------
# Loading
# --------------------------------------------------------------------------


def display_name(raw: str) -> str:
    """``"carl af wirsén"`` -> ``"Carl af Wirsén"`` (particles stay lowercase)."""
    words = raw.split()
    out = []
    for i, w in enumerate(words):
        if i > 0 and w.lower() in _NAME_PARTICLES:
            out.append(w.lower())
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out)


def _iso_date(date_int: int) -> str:
    s = str(date_int)
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}"


def _order_chain(rows: list[tuple]) -> list[tuple]:
    """Order a discussion's rows by walking prev/next inside the group.

    ``rows`` are ``(id, prev, next, *rest)``. Raises ``ValueError`` if the
    chain does not visit every member exactly once.
    """
    by_id = {r[0]: r for r in rows}
    starts = [r for r in rows if r[1] not in by_id]
    if len(starts) != 1:
        raise ValueError(
            f"discussion chain has {len(starts)} start nodes, expected 1: "
            f"{sorted(by_id)}"
        )
    ordered = []
    cur = starts[0]
    seen = set()
    while cur is not None:
        if cur[0] in seen:
            raise ValueError(f"discussion chain loops at {cur[0]}")
        seen.add(cur[0])
        ordered.append(cur)
        cur = by_id.get(cur[2])
    if len(ordered) != len(rows):
        raise ValueError(
            f"discussion chain covers {len(ordered)} of {len(rows)} utterances: "
            f"{sorted(by_id)}"
        )
    return ordered


def load_discussions(conn: sqlite3.Connection, names: dict[str, str]) -> list[Discussion]:
    """Read every discussion from a reduced ``ToK_data`` database.

    ``names`` maps swerik person ids (``person.name`` in the DB) to
    human-readable names; unmapped ids are shown verbatim.
    """
    sql = """
        SELECT u.discussion_id, u.id, u.prev, u.next, p.name, p.gender, p.party,
               u.year, u.date, u.kammare,
               COALESCE(u.kvinna_1, 0), COALESCE(u.kvinna_2, 0), COALESCE(u.kvinna_3, 0),
               u.content
        FROM utterance u
        JOIN person p ON p.id = u.person_id
        WHERE u.discussion_id IS NOT NULL
        ORDER BY u.discussion_id, u.id
    """
    groups: dict[int, list[tuple]] = {}
    for row in conn.execute(sql):
        groups.setdefault(row[0], []).append(row[1:])

    discussions = []
    for disc_id in sorted(groups):
        ordered = _order_chain(groups[disc_id])
        first = ordered[0]
        year, date_int, chamber = first[6], first[7], first[8]
        utterances = [
            Utterance(
                id=r[0],
                order=i + 1,
                speaker_id=r[3],
                speaker_name=display_name(names[r[3]]) if r[3] in names else r[3],
                gender=r[4],
                party=r[5],
                kvinna_1=bool(r[9]),
                kvinna_2=bool(r[10]),
                kvinna_3=bool(r[11]),
                content=r[12] or "",
            )
            for i, r in enumerate(ordered)
        ]
        discussions.append(
            Discussion(
                id=disc_id,
                year=year,
                date=_iso_date(date_int),
                chamber=chamber,
                chamber_name=CHAMBER_NAMES.get(chamber, str(chamber)),
                utterances=utterances,
            )
        )
    return discussions


def load_name_map() -> dict[str, str]:
    """``who -> name`` from the shipped ``speaker_totals.csv.gz``."""
    from .data import read_speaker_totals

    df = read_speaker_totals()
    return {
        str(who): str(name)
        for who, name in zip(df["who"], df["name"])
        if isinstance(name, str) and name.strip()
    }


def iter_all_discussions(years: Iterable[int] | None = None) -> Iterator[Discussion]:
    """Yield discussions from the shipped yearly databases, oldest year first."""
    from .data import _iter_yearly_dbs

    names = load_name_map()
    wanted = set(years) if years is not None else None
    for year, db_path in _iter_yearly_dbs():
        if wanted is not None and year not in wanted:
            continue
        with contextlib.closing(
            sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
        ) as conn:
            yield from load_discussions(conn, names)


# --------------------------------------------------------------------------
# Records and naming
# --------------------------------------------------------------------------


def split_runs(content: str) -> list[tuple[str, bool]]:
    """Split annotated text into ``(text, is_match)`` runs.

    Words carrying ``¹²³`` superscripts (added by ``reduce_db.annotate_content``)
    become bold runs and keep their markers; everything else is plain.
    """
    runs: list[tuple[str, bool]] = []
    pos = 0
    for m in _RUN_RE.finditer(content):
        if m.start() > pos:
            runs.append((content[pos : m.start()], False))
        runs.append((m.group(0), True))
        pos = m.end()
    if pos < len(content):
        runs.append((content[pos:], False))
    return runs


def discussion_record(d: Discussion) -> dict:
    """One aggregate-jsonl line: discussion metadata with nested utterances."""
    return {
        "discussion_id": d.id,
        "year": d.year,
        "date": d.date,
        "chamber": d.chamber,
        "chamber_name": d.chamber_name,
        "n_utterances": len(d.utterances),
        "utterances": [asdict(u) for u in d.utterances],
    }


def utterance_records(d: Discussion) -> list[dict]:
    """Per-discussion jsonl lines: one per utterance, with discussion keys."""
    return [
        {"discussion_id": d.id, "date": d.date, "chamber": d.chamber, **asdict(u)}
        for u in d.utterances
    ]


def file_stem(d: Discussion) -> str:
    return f"{d.date}_{CHAMBER_CODES.get(d.chamber, d.chamber)}_d{d.id:04d}"


def _speaker_heading(u: Utterance) -> str:
    return f"{u.order}. {u.speaker_name} ({u.party or '–'}, {u.gender or '–'})"


def _metadata_lines(d: Discussion) -> list[tuple[str, str]]:
    return [
        (LABELS["date"], d.date),
        (LABELS["chamber"], d.chamber_name),
        (LABELS["n_utterances"], str(len(d.utterances))),
        (LABELS["source"], SOURCE_LINE),
    ]


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def render_docx(d: Discussion) -> Document:
    doc = Document()
    props = doc.core_properties
    props.created = FIXED_DATETIME
    props.modified = FIXED_DATETIME
    props.last_printed = FIXED_DATETIME
    props.author = "Tal om Kvinnor"
    props.title = LABELS["title"].format(id=d.id)

    doc.add_heading(LABELS["title"].format(id=d.id), level=0)
    for label, value in _metadata_lines(d):
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)
    doc.add_paragraph(LABELS["legend"]).runs[0].italic = True

    for u in d.utterances:
        doc.add_heading(_speaker_heading(u), level=2)
        p = doc.add_paragraph()
        for text, is_match in split_runs(u.content):
            run = p.add_run(text)
            run.bold = is_match
    return doc


def docx_bytes(d: Discussion) -> bytes:
    """Render to bytes with zip entry timestamps pinned (python-docx stamps
    entries with the wall clock, which would defeat determinism)."""
    raw = io.BytesIO()
    render_docx(d).save(raw)
    raw.seek(0)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(
        out, "w", compression=zipfile.ZIP_DEFLATED
    ) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


class _DiscussionPDF(FPDF):
    def footer(self) -> None:
        self.set_y(-15)
        self.set_font(FONT_FAMILY, size=8)
        self.set_text_color(120)
        self.cell(0, 8, LABELS["page"].format(n=self.page_no()), align="C")
        self.set_text_color(0)


def render_pdf(d: Discussion) -> bytes:
    pdf = _DiscussionPDF()
    pdf.set_creation_date(FIXED_DATETIME)
    pdf.set_title(LABELS["title"].format(id=d.id))
    pdf.set_author("Tal om Kvinnor")
    pdf.add_font(FONT_FAMILY, "", FONT_DIR / "DejaVuSans.ttf")
    pdf.add_font(FONT_FAMILY, "B", FONT_DIR / "DejaVuSans-Bold.ttf")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=20)
    pdf.add_page()

    line = 6
    pdf.set_font(FONT_FAMILY, "B", 16)
    pdf.cell(0, 10, LABELS["title"].format(id=d.id), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(FONT_FAMILY, size=10)
    for label, value in _metadata_lines(d):
        pdf.set_font(FONT_FAMILY, "B", 10)
        pdf.write(line, f"{label}: ")
        pdf.set_font(FONT_FAMILY, "", 10)
        pdf.write(line, value)
        pdf.ln(line)
    pdf.ln(2)
    pdf.set_font(FONT_FAMILY, "", 9)
    pdf.set_text_color(90)
    pdf.write(line, LABELS["legend"])
    pdf.set_text_color(0)
    pdf.ln(line + 4)

    for u in d.utterances:
        pdf.set_font(FONT_FAMILY, "B", 11)
        pdf.write(line, _speaker_heading(u))
        pdf.ln(line + 1)
        for text, is_match in split_runs(u.content):
            pdf.set_font(FONT_FAMILY, "B" if is_match else "", 10)
            pdf.write(line, text)
        pdf.ln(line + 4)

    return bytes(pdf.output())


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------


def _dump(record: dict) -> str:
    return json.dumps(record, ensure_ascii=False, sort_keys=True)


def export(out_dir: Path | None = None, years: Iterable[int] | None = None) -> int:
    """Write all per-discussion files plus ``discussions.jsonl.gz``.

    Returns the number of discussions written. When ``years`` is given the
    aggregate file only covers those years.
    """
    out_dir = out_dir if out_dir is not None else settings.discussions_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    aggregate = out_dir / "discussions.jsonl.gz"

    count = 0
    with gzip.GzipFile(aggregate, "wb", mtime=0) as agg:
        for d in tqdm(iter_all_discussions(years), desc="Exporting discussions"):
            year_dir = out_dir / str(d.year)
            year_dir.mkdir(exist_ok=True)
            stem = year_dir / file_stem(d)
            stem.with_suffix(".pdf").write_bytes(render_pdf(d))
            stem.with_suffix(".docx").write_bytes(docx_bytes(d))
            stem.with_suffix(".jsonl").write_text(
                "".join(_dump(r) + "\n" for r in utterance_records(d)), encoding="utf-8"
            )
            agg.write((_dump(discussion_record(d)) + "\n").encode("utf-8"))
            count += 1
    return count


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=None, help="output directory")
    parser.add_argument("--years", type=int, nargs="*", default=None, help="restrict to these years")
    args = parser.parse_args(argv)
    n = export(args.out, args.years)
    print(f"Wrote {n:,} discussions to {args.out or settings.discussions_dir}.")


if __name__ == "__main__":
    main()
